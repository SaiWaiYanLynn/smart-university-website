"""Small slide engine. One element list per slide, two renderers (HTML for the PDF, PPTX for editing).

Canvas is 1920 x 1080 px. In the PPTX 1 px = 0.5 pt, so 28 px text is 14 pt.
"""
import os, math, html, json
from PIL import Image, ImageFont

W, H = 1920, 1080
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
IMG = os.path.join(ROOT, 'img')
FONTDIR = os.path.join(ROOT, 'fonts')

# ---- tokens (same names and values as the design system) ----
C = dict(
    teal='#0F766E', teal_dark='#115E59', teal_tint='#F0FDFA', teal_100='#CCFBF1', teal_200='#99F6E4',
    charcoal='#27272A', ink='#18181B', text='#27272A', muted='#52525B', subtle='#71717A',
    border='#E4E4E7', border_strong='#A1A1AA', bg='#FFFFFF', alt='#FAFAFA', panel='#F4F4F5',
    lime='#A3E635', white='#FFFFFF',
    red='#B91C1C', red_tint='#FEF2F2', red_200='#FECACA',
    green='#15803D', green_tint='#F0FDF4', green_200='#BBF7D0',
    amber='#B45309', amber_tint='#FFFBEB', amber_200='#FDE68A',
    blue='#1D4ED8', blue_tint='#EFF6FF',
)

BASE_C = dict(C)
THEMES = {
    'A': dict(),
    'B': dict(bg='#F1F1F3', alt='#FFFFFF', panel='#E4E4E7', border='#D4D4D8', border_strong='#9A9AA3'),
    'C': dict(bg='#FBF8F2', alt='#FFFFFF', panel='#F2EDE2', border='#E3DAC8', border_strong='#B5AA92',
              text='#2A2723', ink='#1C1A17', muted='#5A544A', subtle='#766F63'),
}


def set_theme(name):
    """Swap the color tokens in place. Every module reads C at call time, so a deck takes its theme from this call."""
    C.clear()
    C.update(BASE_C)
    C.update(THEMES[name])
    return name


WEIGHT_FILE = {400: 'Inter-Regular.ttf', 500: 'Inter-Medium.ttf', 600: 'Inter-SemiBold.ttf', 700: 'Inter-Bold.ttf'}
_font_cache = {}


def font(size, weight=400, italic=False):
    key = (round(size * 4), weight, italic)
    if key not in _font_cache:
        fn = 'Inter-Italic.ttf' if italic else WEIGHT_FILE[weight]
        _font_cache[key] = ImageFont.truetype(os.path.join(FONTDIR, fn), size=max(1, int(round(size * 4))))
    return _font_cache[key]


def text_w(s, size, weight=400, italic=False):
    return font(size, weight, italic).getlength(s) / 4.0


def wrap(s, width, size, weight=400, italic=False):
    """Greedy wrap. Returns list of lines. Honors explicit newlines."""
    lines = []
    for para in s.split('\n'):
        words = para.split(' ')
        cur = ''
        for w_ in words:
            trial = (cur + ' ' + w_).strip() if cur else w_
            if text_w(trial, size, weight, italic) <= width or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w_
        lines.append(cur)
    return lines


class Slide:
    def __init__(self, deck, kind='content', bg=None, notes='', number=True, footer=True, section=None):
        self.deck, self.kind, self.notes = deck, kind, notes
        self.section = section
        self.bg = bg or C['bg']
        self.els = []
        self.number = number
        self.footer = footer
        self.issues = []
        self.title_text = ''

    # ---------- primitives ----------
    def rect(self, x, y, w, h, fill=None, stroke=None, sw=2, r=0, shadow=False, name=''):
        self.els.append(dict(t='rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, r=r, shadow=shadow, name=name))

    def ellipse(self, x, y, w, h, fill=None, stroke=None, sw=2):
        self.els.append(dict(t='ellipse', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw))

    def line(self, x1, y1, x2, y2, color=None, w=3, arrow=False, dash=False):
        self.els.append(dict(t='line', x1=x1, y1=y1, x2=x2, y2=y2, color=color or C['border_strong'], w=w, arrow=arrow, dash=dash))

    def text(self, x, y, w, s, size=30, weight=400, color=None, align='l', lh=1.28, italic=False, h=None,
             anchor='t', upper=False, tracking=0, name='', fit=True, minsize=22, runs=None):
        """Place text. Height is measured. If h is given and text overflows, the font shrinks to minsize."""
        color = color or C['text']
        if upper:
            s = s.upper()
        sz = size
        while True:
            lines = wrap(s, w * 0.975, sz, weight, italic)
            hh = len(lines) * sz * lh
            if h is None or hh <= h + 0.5 or not fit or sz <= minsize:
                break
            sz -= 1
        if h is not None and hh > h + 0.5:
            self.issues.append(f'text overflow: "{s[:40]}" needs {hh:.0f} has {h}')
        self.els.append(dict(t='text', x=x, y=y, w=w, h=(h if h is not None else hh), lines=lines, size=sz, weight=weight,
                             color=color, align=align, lh=lh, italic=italic, anchor=anchor, tracking=tracking, name=name,
                             need=hh, raw=s))
        return y + hh

    def image(self, path, x, y, w, h, fit='contain', r=0, border=None, shadow=False, pad=0, bgfill=None):
        full = path if os.path.isabs(path) else os.path.join(IMG, path)
        if not os.path.exists(full):
            self.issues.append('missing image ' + path)
        self.els.append(dict(t='image', path=full, x=x, y=y, w=w, h=h, fit=fit, r=r, border=border, shadow=shadow, pad=pad, bgfill=bgfill))

    # ---------- composed helpers ----------
    def title(self, text, overline=None, color=None, size=54, y=64, w=None):
        self.title_text = text
        th = getattr(self.deck, 'theme', 'A')
        if th == 'B':
            # charcoal band, lime tick under it, white title
            top = 30
            ty = top + (36 if overline else 8)
            if overline:
                self.text(96, top, 1500, overline, size=20, weight=700, color=C['lime'], upper=True, tracking=2.5)
            nlines = len(wrap(text, (w or 1728) * 0.975, min(size, 48), 700))
            tsz = min(size, 48)
            band = ty + nlines * tsz * 1.15 + 30
            self.els.insert(0, dict(t='rect', x=0, y=0, w=W, h=band, fill=C['charcoal'], stroke=None, sw=2, r=0, shadow=False, name='band'))
            self.els.insert(1, dict(t='rect', x=96, y=band - 7, w=170, h=7, fill=C['lime'], stroke=None, sw=2, r=0, shadow=False, name='tick'))
            self.text(96, ty, w or 1728, text, size=tsz, weight=700, color=C['white'], lh=1.15)
            return band - 8
        if th == 'C':
            top = 54
            ty = top + (38 if overline else 0)
            if overline:
                self.text(124, top, 1500, overline, size=22, weight=600, color=C['teal'])
            tsz = min(size + 4, 58)
            nlines = len(wrap(text, ((w or 1728) - 28) * 0.975, tsz, 700))
            hh = nlines * tsz * 1.12
            self.rect(96, top + 2, 8, (ty - top) + hh - 2, fill=C['teal'], r=4)
            end = self.text(124, ty, (w or 1728) - 28, text, size=tsz, weight=700, color=color or C['ink'], lh=1.12)
            return end
        yy = y
        if overline:
            self.text(96, yy, 1100, overline, size=22, weight=600, color=C['teal'], upper=True, tracking=2)
            yy += 40
        end = self.text(96, yy, w or 1728, text, size=size, weight=700, color=color or C['ink'], lh=1.15)
        return end

    def bullets(self, x, y, w, items, size=30, gap=16, color=None, dot=None, lh=1.3, weight=400):
        """items: list of str or (bold_lead, rest). Returns bottom y."""
        yy = y
        for it in items:
            lead, rest = (it if isinstance(it, tuple) else (None, it))
            self.ellipse(x, yy + size * 0.5 - 6, 12, 12, fill=dot or C['teal'])
            if lead:
                end = self.rich(x + 34, yy, w - 34, [(lead + ' ', 700), (rest, weight)], size=size, color=color, lh=lh)
            else:
                end = self.text(x + 34, yy, w - 34, rest, size=size, weight=weight, color=color, lh=lh)
            yy = end + gap
        return yy - gap

    def rich(self, x, y, w, runs, size=30, color=None, lh=1.3, align='l'):
        """Inline bold. runs: list of (text, weight[, color]). Wraps across runs."""
        color = color or C['text']
        # tokenise to words with weights
        words = []
        for r in runs:
            t, wt = r[0], r[1]
            col = r[2] if len(r) > 2 else color
            for i, part in enumerate(t.split(' ')):
                if part == '' and i > 0:
                    continue
                words.append((part, wt, col))
        lines, cur, curw = [], [], 0.0
        sp = text_w(' ', size, 400)
        maxw = w * 0.975
        for word, wt, col in words:
            ww = text_w(word, size, wt)
            add = ww + (sp if cur else 0)
            if cur and curw + add > maxw:
                lines.append(cur)
                cur, curw = [(word, wt, col)], ww
            else:
                cur.append((word, wt, col))
                curw += add
        if cur:
            lines.append(cur)
        hh = len(lines) * size * lh
        self.els.append(dict(t='rich', x=x, y=y, w=w, h=hh, lines=lines, size=size, color=color, lh=lh, align=align, need=hh))
        return y + hh

    def card(self, x, y, w, h, title=None, body=None, fill=None, stroke=None, tsize=30, bsize=26, accent=None, r=16, pad=28,
             tcolor=None, bcolor=None, bullets=None):
        self.rect(x, y, w, h, fill=fill or C['bg'], stroke=stroke or C['border'], sw=2, r=r)
        if accent:
            self.rect(x, y + 18, 8, h - 36, fill=accent, r=4)
        yy = y + pad
        if title:
            yy = self.text(x + pad, yy, w - 2 * pad, title, size=tsize, weight=700, color=tcolor or C['ink'], lh=1.2) + 10
        if body:
            yy = self.text(x + pad, yy, w - 2 * pad, body, size=bsize, color=bcolor or C['muted'], lh=1.3, h=max(10, y + h - pad - yy), minsize=22)
        if bullets:
            yy = self.bullets(x + pad, yy, w - 2 * pad, bullets, size=bsize, gap=10)
        return yy

    def pill(self, x, y, s, fill=None, color=None, size=22, weight=600, padx=16, h=40, stroke=None):
        tw = text_w(s, size, weight)
        self.rect(x, y, tw + 2 * padx, h, fill=fill or C['teal_tint'], stroke=stroke, r=h / 2)
        self.text(x + padx, y + (h - size * 1.2) / 2, tw + 40, s, size=size, weight=weight, color=color or C['teal'], lh=1.2, fit=False)
        return x + tw + 2 * padx

    def mark(self, kind, cx, cy, s=26):
        """Drawn check, cross and partial marks. No glyphs, so they look the same everywhere."""
        if kind == 'y':
            self.ellipse(cx - s / 2, cy - s / 2, s, s, fill=C['green_tint'], stroke=C['green'], sw=2)
            self.line(cx - s * .24, cy + s * .02, cx - s * .06, cy + s * .2, color=C['green'], w=4)
            self.line(cx - s * .06, cy + s * .2, cx + s * .26, cy - s * .18, color=C['green'], w=4)
        elif kind == 'n':
            self.ellipse(cx - s / 2, cy - s / 2, s, s, fill=C['red_tint'], stroke=C['red'], sw=2)
            self.line(cx - s * .2, cy - s * .2, cx + s * .2, cy + s * .2, color=C['red'], w=4)
            self.line(cx - s * .2, cy + s * .2, cx + s * .2, cy - s * .2, color=C['red'], w=4)
        else:
            self.ellipse(cx - s / 2, cy - s / 2, s, s, fill=C['amber_tint'], stroke=C['amber'], sw=2)
            self.line(cx - s * .22, cy, cx + s * .22, cy, color=C['amber'], w=4)

    def device(self, path, x, y, w, caption=None, border=True, shadow=True, csize=22):
        """Frame screenshot in a rounded card. Height follows the image aspect. Returns bottom y."""
        full = os.path.join(IMG, path)
        im = Image.open(full)
        h = w * im.size[1] / im.size[0]
        self.image(path, x, y, w, h, fit='contain', r=14, border=C['border'] if border else None, shadow=shadow)
        yy = y + h
        if caption:
            yy = self.text(x, yy + 12, w, caption, size=csize, color=C['muted'], align='c', lh=1.25)
        return yy

    def callout(self, x, y, w, n, title, body, size=26):
        self.ellipse(x, y, 46, 46, fill=C['teal'])
        self.text(x, y + 7, 46, str(n), size=26, weight=700, color=C['white'], align='c', lh=1.1)
        yy = self.text(x + 66, y + 2, w - 66, title, size=size + 2, weight=700, color=C['ink'], lh=1.2)
        yy = self.text(x + 66, yy + 4, w - 66, body, size=size, color=C['muted'], lh=1.3)
        return yy

    def table(self, x, y, colw, rows, header=True, size=26, hsize=22, pad=18, zebra=True, hfill=None, minrow=56,
              first_bold=True, colcolors=None, cellfill=None, valign='m'):
        """rows: list of list of str. colw: list of widths. Heights come from measured text. Returns bottom y."""
        tw = sum(colw)
        yy = y
        for ri, row in enumerate(rows):
            is_h = header and ri == 0
            sz = hsize if is_h else size
            hts = []
            for ci, cell in enumerate(row):
                wt = 700 if (is_h or (first_bold and ci == 0)) else 400
                hts.append(len(wrap(cell, (colw[ci] - 2 * pad) * 0.975, sz, wt)) * sz * 1.28)
            rh = max(minrow, max(hts) + 2 * (pad * 0.6))
            if is_h:
                self.rect(x, yy, tw, rh, fill=hfill or C['charcoal'], r=0)
            elif zebra and ri % 2 == 0:
                self.rect(x, yy, tw, rh, fill=C['alt'])
            if not is_h:
                self.line(x, yy + rh, x + tw, yy + rh, color=C['border'], w=2)
            cx = x
            for ci, cell in enumerate(row):
                wt = 700 if (is_h or (first_bold and ci == 0)) else 400
                col = C['white'] if is_h else (colcolors[ci] if colcolors and colcolors[ci] else (C['ink'] if wt == 700 else C['text']))
                if cellfill and not is_h and cellfill.get((ri, ci)):
                    self.rect(cx, yy, colw[ci], rh, fill=cellfill[(ri, ci)])
                ty = yy + (rh - hts[ci]) / 2 if valign == 'm' else yy + pad * 0.6
                self.text(cx + pad, ty, colw[ci] - 2 * pad, cell, size=sz, weight=wt, color=col, upper=False, lh=1.28, fit=False)
                cx += colw[ci]
            yy += rh
        self.rect(x, y, tw, yy - y, stroke=C['border'], sw=2, r=0)
        return yy

    # ---------- chrome ----------
    def _tracker(self, x, y, w, size=20, on='#FFFFFF', off='#A1A1AA', hot=None, sep=None):
        """Section tracker: every section name in one line, the current one bold. Right aligned in the box."""
        secs = getattr(self.deck, 'sections', None)
        if not secs or self.section is None:
            return
        hot = hot or C['teal']
        sep = sep or '  ·  '
        parts = []
        for i, nm in enumerate(secs):
            cur = (nm == self.section)
            parts.append((nm, 700 if cur else 400, hot if cur else off))
        total = sum(text_w(p[0], size, p[1]) for p in parts) + text_w(sep, size, 400) * (len(parts) - 1)
        if total > w:
            size = size - 2
            total = sum(text_w(p[0], size, p[1]) for p in parts) + text_w(sep, size, 400) * (len(parts) - 1)
        cx = x + w - total
        for i, (nm, wt, col) in enumerate(parts):
            self.text(cx, y, text_w(nm, size, wt) + 6, nm, size=size, weight=wt, color=col, lh=1.2, fit=False)
            cx += text_w(nm, size, wt)
            if i < len(parts) - 1:
                sw_ = text_w(sep, size, 400)
                self.text(cx, y, sw_, '\u00b7', size=size, color=off, lh=1.2, fit=False, align='c')
                cx += sw_

    def chrome(self, n, total):
        th = getattr(self.deck, 'theme', 'A')
        label = 'Thales University  ·  Smart University Website  ·  Group S'
        if self.kind == 'dark':
            if self.number and th == 'A':
                self.text(1624, 1038, 200, f'{n}', size=22, weight=600, color='#A1A1AA', align='r', lh=1.2)
            return
        if not self.footer:
            return
        if th == 'B':
            self.rect(0, 1034, W, 46, fill=C['charcoal'])
            self.text(96, 1045, 760, label, size=20, color='#A1A1AA', lh=1.2, fit=False)
            self._tracker(820, 1045, 880, size=20, off='#A1A1AA', hot=C['lime'])
            if self.number:
                self.text(1724, 1044, 100, f'{n}', size=22, weight=700, color=C['lime'], align='r', lh=1.2)
        elif th == 'C':
            self.text(96, 1042, 760, label, size=20, color=C['subtle'], lh=1.2, fit=False)
            self._tracker(820, 1042, 880, size=20, off=C['border_strong'], hot=C['teal'])
            if self.number:
                self.text(1724, 1040, 100, f'{n}', size=22, weight=600, color=C['subtle'], align='r', lh=1.2)
        else:
            self.line(96, 1026, 1824, 1026, color=C['border'], w=2)
            self.text(96, 1038, 760, label, size=20, color=C['subtle'], lh=1.2, fit=False)
            self._tracker(820, 1038, 880, size=20, off=C['border_strong'], hot=C['teal'])
            if self.number:
                self.text(1724, 1038, 100, f'{n}', size=22, weight=600, color=C['subtle'], align='r', lh=1.2)

    # ---------- checks ----------
    def check(self):
        out = list(self.issues)
        texts = [e for e in self.els if e['t'] in ('text', 'rich')]
        for e in self.els:
            if e['t'] in ('rect', 'image', 'text', 'rich'):
                if e['x'] < -1 or e['y'] < -1 or e['x'] + e['w'] > W + 1 or e.get('y', 0) + e.get('h', 0) > H + 1:
                    out.append(f"out of canvas: {e['t']} at {e['x']:.0f},{e['y']:.0f} {e['w']:.0f}x{e.get('h', 0):.0f}")
            if e['t'] == 'text' and e['size'] < 20 and e.get('raw','').strip():
                out.append(f"font too small {e['size']}: {e['raw'][:30]}")
        for i in range(len(texts)):
            a = texts[i]
            for j in range(i + 1, len(texts)):
                b = texts[j]
                ox = min(a['x'] + a['w'], b['x'] + b['w']) - max(a['x'], b['x'])
                oy = min(a['y'] + a['need'], b['y'] + b['need']) - max(a['y'], b['y'])
                if ox > 8 and oy > 4:
                    sa = a.get('raw') or ' '.join(w[0] for w in a['lines'][0])
                    sb = b.get('raw') or ' '.join(w[0] for w in b['lines'][0])
                    out.append(f'text overlap: "{sa[:28]}" x "{sb[:28]}" ({ox:.0f}x{oy:.0f})')
        return out


class Deck:
    def __init__(self, name, title, theme='A', sections=None):
        self.name, self.title = name, title
        self.theme = set_theme(theme)
        self.sections = sections or []
        self.slides = []

    def add(self, **kw):
        s = Slide(self, **kw)
        self.slides.append(s)
        return s

    def report(self):
        n = 0
        for i, s in enumerate(self.slides, 1):
            for m in s.check():
                n += 1
                print(f'  [{self.name} #{i} {s.title_text[:40]}] {m}')
        return n


# ======================= HTML renderer =======================

def _css_font(weight, italic=False):
    return 'italic' if italic else 'normal'


def render_html(deck, path):
    parts = []
    face = ''.join(
        f"@font-face{{font-family:'InterX';font-weight:{wt};src:url('file://{FONTDIR}/{fn}');}}"
        for wt, fn in WEIGHT_FILE.items())
    face += f"@font-face{{font-family:'InterX';font-style:italic;font-weight:400;src:url('file://{FONTDIR}/Inter-Italic.ttf');}}"
    css = face + f"""
@page{{size:{W}px {H}px;margin:0}}
*{{box-sizing:border-box;margin:0;padding:0}}
html,body{{background:#333}}
body{{font-family:'InterX',sans-serif}}
.slide{{position:relative;width:{W}px;height:{H}px;overflow:hidden;page-break-after:always;break-after:page}}
.slide:last-child{{page-break-after:auto;break-after:auto}}
.e{{position:absolute}}
.tx{{white-space:nowrap}}
@media screen{{.slide{{margin:20px auto}}}}
"""
    total = len(deck.slides)
    for n, s in enumerate(deck.slides, 1):
        s_els = list(s.els)
        # chrome is appended last so it sits on top
        chrome = Slide(deck, kind=s.kind)
        chrome.footer, chrome.number, chrome.section = s.footer, s.number, s.section
        chrome.chrome(n, total)
        els = s_els + chrome.els
        out = [f'<section class="slide" style="background:{s.bg}">']
        for e in els:
            out.append(_html_el(e))
        out.append('</section>')
        parts.append('\n'.join(out))
    with open(path, 'w') as f:
        f.write(f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(deck.title)}</title><style>{css}</style></head><body>' + '\n'.join(parts) + '</body></html>')


def _shadow(e):
    return 'box-shadow:0 6px 24px rgba(24,24,27,.14);' if e.get('shadow') else ''


def _html_el(e):
    t = e['t']
    if t == 'rect':
        st = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;"
        if e['fill']:
            st += f"background:{e['fill']};"
        if e['stroke']:
            st += f"border:{e['sw']}px solid {e['stroke']};"
        if e['r']:
            st += f"border-radius:{e['r']}px;"
        st += _shadow(e)
        return f'<div class="e" style="{st}"></div>'
    if t == 'ellipse':
        st = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;border-radius:50%;"
        if e['fill']:
            st += f"background:{e['fill']};"
        if e['stroke']:
            st += f"border:{e['sw']}px solid {e['stroke']};"
        return f'<div class="e" style="{st}"></div>'
    if t == 'line':
        dx, dy = e['x2'] - e['x1'], e['y2'] - e['y1']
        L = math.hypot(dx, dy)
        ang = math.degrees(math.atan2(dy, dx))
        st = (f"left:{e['x1']}px;top:{e['y1'] - e['w'] / 2}px;width:{L}px;height:{e['w']}px;background:{e['color']};"
              f"transform-origin:0 50%;transform:rotate({ang}deg);border-radius:{e['w'] / 2}px;")
        if e['dash']:
            st = (f"left:{e['x1']}px;top:{e['y1'] - e['w'] / 2}px;width:{L}px;height:{e['w']}px;"
                  f"background:repeating-linear-gradient(90deg,{e['color']} 0 14px,transparent 14px 24px);"
                  f"transform-origin:0 50%;transform:rotate({ang}deg);")
        s = f'<div class="e" style="{st}"></div>'
        if e['arrow']:
            ah = 16
            st2 = (f"left:{e['x2']}px;top:{e['y2']}px;width:0;height:0;"
                   f"border-left:{ah}px solid {e['color']};border-top:{ah * .6}px solid transparent;border-bottom:{ah * .6}px solid transparent;"
                   f"transform-origin:0 0;transform:translate({-ah}px,{-ah * .6}px) rotate({ang}deg);")
            # rotate around the tip: place tip at (x2,y2)
            st2 = (f"left:{e['x2']}px;top:{e['y2']}px;width:0;height:0;border-left:{ah}px solid {e['color']};"
                   f"border-top:{ah * .6}px solid transparent;border-bottom:{ah * .6}px solid transparent;"
                   f"transform-origin:{ah}px {ah * .6}px;transform:translate({-ah}px,{-ah * .6}px) rotate({ang}deg);")
            s += f'<div class="e" style="{st2}"></div>'
        return s
    if t == 'image':
        full = e['path']
        if not os.path.exists(full):
            return ''
        im = Image.open(full)
        iw, ih = im.size
        bw, bh = e['w'] - 2 * e['pad'], e['h'] - 2 * e['pad']
        if e['fit'] == 'contain':
            sc = min(bw / iw, bh / ih)
            dw, dh = iw * sc, ih * sc
            dx, dy = e['x'] + e['pad'] + (bw - dw) / 2, e['y'] + e['pad'] + (bh - dh) / 2
            wrap_st = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;"
            if e.get('bgfill'):
                wrap_st += f"background:{e['bgfill']};"
            if e['border']:
                wrap_st += f"border:2px solid {e['border']};"
            if e['r']:
                wrap_st += f"border-radius:{e['r']}px;"
            wrap_st += _shadow(e)
            img = (f"<img src='file://{full}' style='position:absolute;left:{dx - e['x'] - (2 if e['border'] else 0)}px;top:{dy - e['y'] - (2 if e['border'] else 0)}px;"
                   f"width:{dw}px;height:{dh}px;" + (f"border-radius:{max(0, e['r'] - 2)}px;" if e['r'] else '') + "'>")
            return f'<div class="e" style="{wrap_st}overflow:hidden">{img}</div>'
        # cover
        st = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;overflow:hidden;"
        if e['r']:
            st += f"border-radius:{e['r']}px;"
        if e['border']:
            st += f"border:2px solid {e['border']};"
        st += _shadow(e)
        return f"<div class='e' style=\"{st}\"><img src='file://{full}' style='width:100%;height:100%;object-fit:cover;object-position:top'></div>"
    if t == 'text':
        out = []
        lh_px = e['size'] * e['lh']
        total_h = len(e['lines']) * lh_px
        y0 = e['y']
        if e['anchor'] == 'm':
            y0 = e['y'] + (e['h'] - total_h) / 2
        for i, ln in enumerate(e['lines']):
            st = (f"left:{e['x']}px;top:{y0 + i * lh_px}px;width:{e['w']}px;height:{lh_px}px;line-height:{lh_px}px;"
                  f"font-size:{e['size']}px;font-weight:{e['weight']};color:{e['color']};text-align:{ {'l': 'left', 'c': 'center', 'r': 'right'}[e['align']] };")
            if e['italic']:
                st += 'font-style:italic;'
            if e['tracking']:
                st += f"letter-spacing:{e['tracking']}px;"
            out.append(f'<div class="e tx" style="{st}">{html.escape(ln)}</div>')
        return ''.join(out)
    if t == 'rich':
        out = []
        lh_px = e['size'] * e['lh']
        for i, ln in enumerate(e['lines']):
            spans = ''
            for k, (word, wt, col) in enumerate(ln):
                spans += f'<span style="font-weight:{wt};color:{col}">{html.escape(word)}</span>' + ('&nbsp;' if k < len(ln) - 1 else '')
            st = (f"left:{e['x']}px;top:{e['y'] + i * lh_px}px;width:{e['w']}px;height:{lh_px}px;line-height:{lh_px}px;font-size:{e['size']}px;"
                  f"text-align:{ {'l': 'left', 'c': 'center', 'r': 'right'}[e['align']] };")
            out.append(f'<div class="e tx" style="{st}">{spans}</div>')
        return ''.join(out)
    return ''


# ======================= PPTX renderer =======================
PX = 6350  # EMU per px


def _rgb(hexs):
    from pptx.dml.color import RGBColor
    h = hexs.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _font_name(weight, italic):
    return {400: 'Inter', 500: 'Inter Medium', 600: 'Inter SemiBold', 700: 'Inter'}[weight]


def render_pptx(deck, path):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.oxml.ns import qn
    from lxml import etree

    prs = Presentation()
    prs.slide_width = Emu(W * PX)
    prs.slide_height = Emu(H * PX)
    blank = prs.slide_layouts[6]
    total = len(deck.slides)

    def emu(v):
        return Emu(int(round(v * PX)))

    def style_shape(shp, e):
        if e.get('fill'):
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(e['fill'])
        else:
            shp.fill.background()
        if e.get('stroke'):
            shp.line.color.rgb = _rgb(e['stroke'])
            shp.line.width = Emu(int(e.get('sw', 2) * PX))
        else:
            shp.line.fill.background()
        # no theme style, no shadow
        st = shp._element.find(qn('p:style'))
        if st is not None:
            shp._element.remove(st)
        sp = shp._element.spPr
        ef = etree.SubElement(sp, qn('a:effectLst'))

    def add_text(sl, e):
        x, y, w, h = e['x'], e['y'], e['w'], e['h']
        size, weight, color, align, lh, italic, anchor = e['size'], e['weight'], e['color'], e['align'], e['lh'], e['italic'], e['anchor']
        tb = sl.shapes.add_textbox(emu(x), emu(y), emu(w), emu(max(h, size * lh)))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == 'm' else MSO_ANCHOR.TOP
        for i, para in enumerate(e['raw'].split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
            p.line_spacing = lh / 1.21
            r = p.add_run()
            r.text = para
            f = r.font
            f.size = Pt(size / 2)
            f.name = _font_name(weight, italic)
            f.bold = (weight == 700)
            f.italic = italic
            f.color.rgb = _rgb(color)
            if e.get('tracking'):
                r._r.get_or_add_rPr().set('spc', str(int(e['tracking'] * 50)))
        return tb

    for n, s in enumerate(deck.slides, 1):
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(s.bg)
        chrome = Slide(deck, kind=s.kind)
        chrome.footer, chrome.number, chrome.section = s.footer, s.number, s.section
        chrome.chrome(n, total)
        for e in s.els + chrome.els:
            t = e['t']
            if t == 'rect':
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['r'] else MSO_SHAPE.RECTANGLE, emu(e['x']), emu(e['y']), emu(e['w']), emu(e['h']))
                if e['r']:
                    shp.adjustments[0] = min(0.5, e['r'] / max(1, min(e['w'], e['h'])))
                style_shape(shp, e)
                if e.get('shadow'):
                    # soft outer shadow
                    sp = shp._element.spPr
                    eff = sp.find(qn('a:effectLst'))
                    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad='190500', dist='63500', dir='5400000', algn='t', rotWithShape='0')
                    clr = etree.SubElement(sh, qn('a:srgbClr'), val='18181B')
                    etree.SubElement(clr, qn('a:alpha'), val='16000')
            elif t == 'ellipse':
                shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, emu(e['x']), emu(e['y']), emu(e['w']), emu(e['h']))
                style_shape(shp, e)
            elif t == 'line':
                con = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(e['x1']), emu(e['y1']), emu(e['x2']), emu(e['y2']))
                cst = con._element.find(qn('p:style'))
                if cst is not None:
                    con._element.remove(cst)
                con.line.color.rgb = _rgb(e['color'])
                con.line.width = Emu(int(e['w'] * PX))
                ln = con.line._get_or_add_ln()
                if e['dash']:
                    etree.SubElement(ln, qn('a:prstDash'), val='dash')
                if e['arrow']:
                    etree.SubElement(ln, qn('a:tailEnd'), type='triangle')
                if e['w'] <= 4:
                    ln.set('cap', 'rnd')
            elif t == 'image':
                full = e['path']
                if not os.path.exists(full):
                    continue
                im = Image.open(full)
                iw, ih = im.size
                bw, bh = e['w'] - 2 * e['pad'], e['h'] - 2 * e['pad']
                if e['border'] or e.get('bgfill') or e['shadow']:
                    box = {'x': e['x'], 'y': e['y'], 'w': e['w'], 'h': e['h'], 'fill': e.get('bgfill'), 'stroke': e['border'], 'sw': 2, 'r': e['r']}
                    shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['r'] else MSO_SHAPE.RECTANGLE, emu(e['x']), emu(e['y']), emu(e['w']), emu(e['h']))
                    if e['r']:
                        shp.adjustments[0] = min(0.5, e['r'] / max(1, min(e['w'], e['h'])))
                    style_shape(shp, box)
                    if e['shadow']:
                        sp = shp._element.spPr
                        eff = sp.find(qn('a:effectLst'))
                        sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad='190500', dist='63500', dir='5400000', algn='t', rotWithShape='0')
                        clr = etree.SubElement(sh, qn('a:srgbClr'), val='18181B')
                        etree.SubElement(clr, qn('a:alpha'), val='16000')
                inset = 2 if e['border'] else 0
                if e['fit'] == 'contain':
                    sc = min((bw - 2 * inset) / iw, (bh - 2 * inset) / ih)
                    dw, dh = iw * sc, ih * sc
                    pic = sl.shapes.add_picture(full, emu(e['x'] + e['pad'] + inset + (bw - 2 * inset - dw) / 2), emu(e['y'] + e['pad'] + inset + (bh - 2 * inset - dh) / 2), emu(dw), emu(dh))
                else:
                    sc = max(bw / iw, bh / ih)
                    dw, dh = iw * sc, ih * sc
                    pic = sl.shapes.add_picture(full, emu(e['x']), emu(e['y']), emu(bw), emu(bh))
                    # crop to top-aligned cover
                    cx = (dw - bw) / 2 / dw
                    pic.crop_left = cx
                    pic.crop_right = cx
                    pic.crop_top = 0
                    pic.crop_bottom = (dh - bh) / dh
            elif t == 'text':
                add_text(sl, e)
            elif t == 'rich':
                tb = sl.shapes.add_textbox(emu(e['x']), emu(e['y']), emu(e['w']), emu(e['h']))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[e['align']]
                p.line_spacing = e['lh'] / 1.21
                first = True
                for ln in e['lines']:
                    for k, (word, wt, col) in enumerate(ln):
                        r = p.add_run()
                        r.text = word + (' ' if True else '')
                        f = r.font
                        f.size = Pt(e['size'] / 2)
                        f.name = _font_name(wt, False)
                        f.bold = (wt == 700)
                        f.color.rgb = _rgb(col)
        # one text per element is rebuilt for plain text above; notes
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
    prs.save(path)
